"""Tests for deployment date range and days deployed.

These figures go into a wildlife exemption filing, so the counting convention is
pinned down here deliberately: days are inclusive of both the first and last day.
"""

import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "critter_counter"))

import pipeline  # noqa: E402
from pipeline import Deployment, Event, ImageResult, deployment_from_events  # noqa: E402


def _image(when: datetime, has_detection: bool = False, species: str = "blank"):
    return ImageResult(
        filepath=Path(f"{when:%Y%m%d%H%M%S}.JPG"),
        timestamp=when,
        has_detection=has_detection,
        species=species,
        confidence=0.9,
        animal_box_count=1 if has_detection else 0,
    )


def test_range_spans_first_to_last_photo():
    events = [
        Event(images=[_image(datetime(2026, 5, 11, 18, 37))]),
        Event(images=[_image(datetime(2026, 6, 21, 13, 30))]),
    ]
    deployment = deployment_from_events(events)

    assert deployment.start == datetime(2026, 5, 11, 18, 37)
    assert deployment.end == datetime(2026, 6, 21, 13, 30)
    assert deployment.date_range == "2026-05-11 to 2026-06-21"


def test_days_are_inclusive_of_both_ends():
    """May 11 to June 21 is 42 days counted inclusively, not 41."""

    events = [
        Event(images=[_image(datetime(2026, 5, 11, 18, 37))]),
        Event(images=[_image(datetime(2026, 6, 21, 13, 30))]),
    ]
    assert deployment_from_events(events).days == 42


def test_single_day_deployment_counts_as_one():
    events = [
        Event(images=[_image(datetime(2026, 5, 11, 6, 0))]),
        Event(images=[_image(datetime(2026, 5, 11, 23, 0))]),
    ]
    assert deployment_from_events(events).days == 1


def test_blank_photos_count_toward_the_deployment_window():
    """Blanks prove the camera was still running; excluding them understates effort."""

    events = [
        Event(images=[_image(datetime(2026, 5, 1), has_detection=False)]),
        Event(images=[_image(datetime(2026, 5, 20), has_detection=True, species="deer")]),
        Event(images=[_image(datetime(2026, 6, 1), has_detection=False)]),
    ]
    deployment = deployment_from_events(events)

    assert deployment.date_range == "2026-05-01 to 2026-06-01"
    assert deployment.days == 32
    assert deployment.photo_count == 3


def test_no_photos_gives_no_range():
    deployment = deployment_from_events([])
    assert deployment.days == 0
    assert deployment.date_range == "no photos"
    assert deployment.photo_count == 0


def _ensemble_file(tmp_path: Path, photos: list[tuple[str, datetime]]) -> Path:
    import json

    from PIL import Image

    predictions = []
    for name, when in photos:
        path = tmp_path / name
        image = Image.new("RGB", (8, 8))
        exif = image.getexif()
        exif.get_ifd(0x8769)[0x9003] = when.strftime("%Y:%m:%d %H:%M:%S")
        image.save(path, exif=exif)
        predictions.append(
            {
                "filepath": str(path),
                "prediction": "id;;;;;;blank",
                "prediction_score": 0.99,
                "detections": [],
            }
        )

    out = tmp_path / "ensemble.json"
    out.write_text(json.dumps({"predictions": predictions}), encoding="utf-8")
    return out


def test_report_can_be_rebuilt_after_the_card_is_removed(tmp_path):
    """Renaming folders and regenerating must not require the card be plugged in."""

    photos = [
        ("a.JPG", datetime(2026, 5, 11, 8, 0)),
        ("b.JPG", datetime(2026, 6, 21, 9, 0)),
    ]
    ensemble = _ensemble_file(tmp_path, photos)
    cache = tmp_path / "timestamps.json"

    first = pipeline.load_results(ensemble, timestamp_cache=cache)
    assert len(first) == 2
    assert cache.exists()

    # Simulate the card being taken out.
    for name, _ in photos:
        (tmp_path / name).unlink()

    second = pipeline.load_results(ensemble, timestamp_cache=cache)
    assert len(second) == 2
    rebuilt = deployment_from_events(pipeline.group_into_events(list(second.values())))
    assert rebuilt.date_range == "2026-05-11 to 2026-06-21"
    assert rebuilt.days == 42


def test_undatable_photos_are_dropped_rather_than_guessed(tmp_path):
    ensemble = _ensemble_file(tmp_path, [("a.JPG", datetime(2026, 5, 11, 8, 0))])
    (tmp_path / "a.JPG").unlink()

    # No cache, no file: nothing can be said about when it was taken.
    assert pipeline.load_results(ensemble, timestamp_cache=tmp_path / "none.json") == {}


def test_report_records_deployment(tmp_path):
    from openpyxl import load_workbook

    events = [
        Event(images=[_image(datetime(2026, 5, 11), has_detection=True, species="deer")]),
        Event(images=[_image(datetime(2026, 6, 21))]),
    ]
    dest = tmp_path / "report.xlsx"
    pipeline.write_excel_report(events, dest, tmp_path)

    rows = list(load_workbook(dest)["Summary"].iter_rows(values_only=True))
    assert rows[0][0] == "Deployment Date Range"
    assert rows[0][1] == "2026-05-11 to 2026-06-21"
    assert rows[1] == ("Count of Deployed (days)", 42, None, None)
    assert rows[2][0] == "Photos Reviewed"
    # Species table still follows, after the blank spacer row.
    assert rows[4][0] == "Species"
    assert any(row[0] == "deer" for row in rows[5:])


def test_combined_report_has_a_deployment_per_folder(tmp_path):
    from openpyxl import load_workbook

    batches = [
        ("North gate", [Event(images=[_image(datetime(2026, 5, 1))]),
                        Event(images=[_image(datetime(2026, 5, 10))])]),
        ("South gate", [Event(images=[_image(datetime(2026, 6, 1))]),
                        Event(images=[_image(datetime(2026, 6, 30))])]),
    ]
    dest = tmp_path / "combined.xlsx"
    pipeline.write_combined_report(batches, dest)

    wb = load_workbook(dest)
    rows = {r[0]: r for r in wb["Deployments"].iter_rows(min_row=2, values_only=True)}

    assert rows["North gate"][1] == "2026-05-01 to 2026-05-10"
    assert rows["North gate"][4] == 10
    assert rows["South gate"][4] == 30

    # The overall range spans every folder.
    overall = list(wb["All folders"].iter_rows(values_only=True))
    assert overall[0][1] == "2026-05-01 to 2026-06-30"
    assert overall[1][1] == 61
