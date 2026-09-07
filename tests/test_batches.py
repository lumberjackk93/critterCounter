"""Tests for splitting a parent folder into separate jobs.

Keeping folders separate is a correctness matter, not tidiness: events are grouped by
how close together photos were taken, so processing two cameras as one batch merges
unrelated sightings whenever their clocks line up.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "critter_counter"))

import pipeline  # noqa: E402
from pipeline import Event, ImageResult, find_batches, folder_image_count  # noqa: E402
from datetime import datetime  # noqa: E402


def _photo(folder: Path, name: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / name
    path.write_bytes(b"jpg")
    return path


def test_each_subfolder_becomes_its_own_job(tmp_path):
    _photo(tmp_path / "DCIM" / "100STLTH", "a.JPG")
    _photo(tmp_path / "New folder", "b.JPG")
    _photo(tmp_path / "New folder (2)", "c.JPG")

    batches = find_batches(tmp_path)
    assert [b.name for b in batches] == ["DCIM", "New folder", "New folder (2)"]


def test_nested_camera_folders_stay_in_one_job(tmp_path):
    """A single card spanning 100XXX/101XXX is one camera, not two jobs."""

    _photo(tmp_path / "DCIM" / "100STLTH", "a.JPG")
    _photo(tmp_path / "DCIM" / "101STLTH", "b.JPG")

    batches = find_batches(tmp_path)
    assert [b.name for b in batches] == ["DCIM"]
    assert folder_image_count(batches[0]) == 2


def test_empty_subfolders_are_skipped(tmp_path):
    _photo(tmp_path / "has_photos", "a.JPG")
    (tmp_path / "empty").mkdir()

    assert [b.name for b in find_batches(tmp_path)] == ["has_photos"]


def test_loose_photos_make_it_a_single_job(tmp_path):
    """A plain camera dump, not a collection of them."""

    _photo(tmp_path, "a.JPG")
    _photo(tmp_path / "thumbnails", "b.JPG")

    assert find_batches(tmp_path) == [tmp_path]


def test_folder_with_no_photos_returns_itself(tmp_path):
    assert find_batches(tmp_path) == [tmp_path]


def _event(species: str, count: int = 1) -> Event:
    images = [
        ImageResult(
            filepath=Path(f"{species}{i}.JPG"),
            timestamp=datetime(2026, 5, 11, 12, 0, i),
            has_detection=True,
            species=species,
            confidence=0.9,
            animal_box_count=1,
        )
        for i in range(count)
    ]
    return Event(images=images)


def test_combined_report_totals_across_folders(tmp_path):
    from openpyxl import load_workbook

    batches = [
        ("North camera", [_event("white-tailed deer"), _event("northern raccoon", 2)]),
        ("South camera", [_event("white-tailed deer", 3)]),
    ]
    dest = tmp_path / "combined.xlsx"
    pipeline.write_combined_report(batches, dest)

    wb = load_workbook(dest)
    rows = {r[0]: r for r in wb["All folders"].iter_rows(min_row=2, values_only=True)}

    # Deer seen in both folders: 2 events, 4 photos.
    assert rows["white-tailed deer"][1] == 2
    assert rows["white-tailed deer"][2] == 4
    assert "North camera" in rows["white-tailed deer"][3]
    assert "South camera" in rows["white-tailed deer"][3]

    assert rows["northern raccoon"][1] == 1
    assert rows["northern raccoon"][2] == 2
    assert rows["northern raccoon"][3] == "North camera"


def test_combined_report_records_folders_with_nothing_found(tmp_path):
    from openpyxl import load_workbook

    dest = tmp_path / "combined.xlsx"
    pipeline.write_combined_report([("Quiet camera", [])], dest)

    wb = load_workbook(dest)
    rows = list(wb["By folder"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "Quiet camera"
    assert rows[0][1] == "(nothing found)"
